import io
import hashlib
import json
import re
import zipfile
from xml.etree import ElementTree

import openpyxl
import xlrd


class ImportErrorDetail(ValueError):
    pass


def _number(value, pattern):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and re.fullmatch(r"0+", pattern or ""):
        return str(value).zfill(len(pattern))
    return str(value)


def _clean(value, pattern=""):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _number(value, pattern)
    return str(value)


def _xls(data):
    warnings = []
    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True, logfile=io.StringIO())
    except xlrd.compdoc.CompDocError:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True, ignore_workbook_corruption=True, logfile=io.StringIO())
        warnings.append("工作簿结构存在异常，已以容错模式读取；请核对预览与源文件")
    sheet = book.sheet_by_index(0)
    grid = []
    for row in range(sheet.nrows):
        values = []
        for col in range(sheet.ncols):
            cell = sheet.cell(row, col)
            if cell.ctype == xlrd.XL_CELL_ERROR:
                raise ImportErrorDetail(f"第 {row + 1} 行第 {col + 1} 列含错误单元格")
            pattern = book.xf_list[cell.xf_index].format_key
            fmt = book.format_map[pattern].format_str if pattern is not None else ""
            values.append(_clean(cell.value, fmt))
        grid.append(values)
    return grid, warnings


def _xlsx(data):
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    sheet = book.worksheets[0]
    grid = [[_clean(cell.value, cell.number_format) for cell in row] for row in sheet.iter_rows()]
    book.close()
    return grid, []


def _xmind(data):
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ImportErrorDetail("XMind 压缩包损坏") from exc
    files = archive.infolist()
    if len(files) > 3000 or sum(item.file_size for item in files) > 20 * 1024 * 1024:
        raise ImportErrorDetail("XMind 解压后内容过大")
    if any(item.compress_size and item.file_size / item.compress_size > 1000 for item in files):
        raise ImportErrorDetail("XMind 压缩比异常")
    titles = []
    if "content.json" in archive.namelist():
        sheets = json.loads(archive.read("content.json"))

        def visit(node, path):
            title = node.get("title", "").strip()
            current = path + ([title] if title else [])
            branches = node.get("children") or {}
            children = branches.get("attached", []) + branches.get("detached", [])
            if children:
                for child in children:
                    visit(child, current)
            elif title:
                titles.append(current)

        for sheet in sheets:
            visit(sheet.get("rootTopic", {}), [])
    elif "content.xml" in archive.namelist():
        root = ElementTree.fromstring(archive.read("content.xml"))

        def child_topics(node):
            for child in node:
                if child.tag.split("}")[-1] == "topic":
                    yield child
                else:
                    yield from child_topics(child)

        def visit(node, path):
            title = next((el.text or "" for el in node if el.tag.split("}")[-1] == "title"), "").strip()
            current = path + ([title] if title else [])
            direct = list(child_topics(node))
            if direct:
                for child in direct:
                    visit(child, current)
            elif title:
                titles.append(current)

        for sheet in root:
            if sheet.tag.split("}")[-1] == "sheet":
                topic = next((el for el in sheet if el.tag.split("}")[-1] == "topic"), None)
                if topic is not None:
                    visit(topic, [])
    else:
        raise ImportErrorDetail("缺少 content.json 或 content.xml")
    if not titles:
        raise ImportErrorDetail("未找到可导入的 XMind 叶节点")
    fingerprint = hashlib.sha256(data).hexdigest()[:10].upper()
    return [["测试用例ID", "用例名称", "用例目录"]] + [
        [f"XM-{fingerprint}-{index:06d}", path[-1], " / ".join(path[:-1])] for index, path in enumerate(titles, 1)
    ], ["XMind 原节点无用例 ID，已生成工作台内部 ID"]


def parse(filename, data):
    suffix = filename.lower().rsplit(".", 1)[-1]
    if len(data) > 10 * 1024 * 1024:
        raise ImportErrorDetail("文件不能超过 10 MB")
    try:
        if suffix == "xls":
            grid, warnings = _xls(data)
        elif suffix == "xlsx":
            grid, warnings = _xlsx(data)
        elif suffix == "xmind":
            grid, warnings = _xmind(data)
        else:
            raise ImportErrorDetail("只支持 .xls、.xlsx 和 .xmind")
    except ImportErrorDetail:
        raise
    except (ValueError, KeyError, IndexError, AttributeError, TypeError, OSError, zipfile.BadZipFile, ElementTree.ParseError, xlrd.XLRDError) as exc:
        raise ImportErrorDetail("文件结构无法解析") from exc
    if not grid:
        raise ImportErrorDetail("文件没有表头")
    headers = [cell or "" for cell in grid[0]]
    if "测试用例ID" not in headers or "用例名称" not in headers:
        raise ImportErrorDetail("缺少测试用例ID或用例名称列")
    rows = [row[:len(headers)] + [None] * max(0, len(headers) - len(row)) for row in grid[1:] if any(cell is not None for cell in row)]
    ids = [row[headers.index("测试用例ID")] for row in rows]
    if any(not value for value in ids) or len(set(ids)) != len(ids):
        raise ImportErrorDetail("用例 ID 为空或在文件内重复")
    if any(not row[headers.index("用例名称")] for row in rows):
        raise ImportErrorDetail("存在空用例名称")
    return headers, rows, warnings
