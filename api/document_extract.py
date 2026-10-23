"""Extract reviewable text without invoking a model or trusting document instructions."""

import io
import os
import signal
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path, PurePosixPath

import pypdfium2 as pdfium
from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from .observability import sanitized_subprocess_environment

ALLOWED = {".md", ".markdown", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".pdf", ".doc", ".docx"}
MAX_DOCX_ENTRIES = 4096
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCX_MEMBER_BYTES = 64 * 1024 * 1024
try:
    MAX_EXTRACTED_TEXT_CHARS = max(
        0, int(os.getenv("AI_STUDIO_MAX_EXTRACTED_TEXT_CHARS", "0"))
    )
except ValueError:
    MAX_EXTRACTED_TEXT_CHARS = 0


class DocumentExtractTimeout(ValueError):
    """Raised when an external document converter exceeds its deadline."""


def _tool_environment() -> dict[str, str]:
    """Keep application and provider credentials out of parser subprocesses."""
    return sanitized_subprocess_environment()


def _stop_external_process(process: subprocess.Popen[bytes]) -> None:
    """Terminate a converter/OCR process and any descendants on timeout."""
    try:
        if os.name != "nt":
            os.killpg(process.pid, signal.SIGTERM)
        else:
            process.terminate()
    except (ProcessLookupError, OSError):
        return
    try:
        process.wait(timeout=2)
    except subprocess.TimeoutExpired:
        try:
            if os.name != "nt":
                os.killpg(process.pid, signal.SIGKILL)
            else:
                process.kill()
        except (ProcessLookupError, OSError):
            pass


def _run_external(command: list[str], *, input_data: bytes | None = None,
                  timeout: float) -> subprocess.CompletedProcess[bytes]:
    """Run a document tool in its own process group with a hard deadline."""
    process = subprocess.Popen(
        command,
        stdin=subprocess.PIPE if input_data is not None else subprocess.DEVNULL,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        shell=False,
        start_new_session=(os.name != "nt"),
        env=_tool_environment(),
    )
    try:
        stdout, stderr = process.communicate(input=input_data, timeout=timeout)
    except subprocess.TimeoutExpired:
        _stop_external_process(process)
        process.communicate()
        raise
    return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)


def _ocr(content: bytes) -> tuple[str, str | None]:
    executable = shutil.which("tesseract")
    if not executable:
        return "", "本机未安装 Tesseract，图片或扫描页无法识别；请安装中文和英文语言包后重新上传"
    try:
        result = _run_external([executable, "stdin", "stdout", "-l", "chi_sim+eng"],
                               input_data=content, timeout=45)
    except subprocess.TimeoutExpired:
        return "", "OCR 超时，请使用更清晰或更小的图片"
    if result.returncode:
        return "", "OCR 失败，请检查 Tesseract 的 chi_sim 与 eng 语言包"
    return result.stdout.decode("utf-8", errors="replace").strip(), None


def _docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    chunks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def _validate_docx_archive(content: bytes) -> None:
    """Reject malformed or highly expanding OOXML archives before parsing.

    The upload limit applies to compressed bytes.  Checking the central
    directory first prevents a small ZIP bomb or path traversal entry from
    expanding into unbounded memory while ``python-docx`` reads it.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise ValueError("Word 文档包含过多压缩文件")
            total = 0
            for info in infos:
                normalized_name = info.filename.replace("\\", "/")
                member = PurePosixPath(normalized_name)
                if "\x00" in info.filename or member.is_absolute() or ".." in member.parts:
                    raise ValueError("Word 文档包含不安全的文件路径")
                if info.file_size > MAX_DOCX_MEMBER_BYTES:
                    raise ValueError("Word 文档内部文件超过大小限制")
                total += info.file_size
                if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ValueError("Word 文档解压后超过大小限制")
    except zipfile.BadZipFile as exc:
        raise ValueError("Word 文档损坏或无法读取") from exc


def extract(filename: str, content: bytes) -> tuple[str, list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError("仅支持 Markdown、文本、图片、PDF 和 Word 文档")
    warnings = []
    if suffix in {".md", ".markdown", ".txt"}:
        text = content.decode("utf-8-sig", errors="replace")
        if "\ufffd" in text:
            warnings.append("部分字符无法按 UTF-8 解码，请核对原文")
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            image = Image.open(io.BytesIO(content))
            if image.width * image.height > 30_000_000:
                raise ValueError("图片像素超过限制")
            image.verify()
        # Pillow raises ``DecompressionBombError`` (an ``Exception`` rather
        # than an ``OSError``) before width/height checks for images that
        # exceed its own safety threshold. Keep that parser failure inside the
        # documented 422 validation path instead of leaking a 500.
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise ValueError("图片损坏或格式不支持") from exc
        text, warning = _ocr(content)
        if warning:
            warnings.append(warning)
    elif suffix == ".doc":
        executable = shutil.which("soffice")
        if not executable:
            raise ValueError("读取旧版 .doc 需要 LibreOffice，请安装后重试")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)
            (path / "source.doc").write_bytes(content)
            try:
                proc = _run_external([executable, f"-env:UserInstallation={(path / 'profile').as_uri()}",
                                      "--headless", "--convert-to", "docx", "--outdir", directory,
                                      str(path / "source.doc")], timeout=90)
            except subprocess.TimeoutExpired as exc:
                # Keep extractor failures in the documented 422 path instead
                # of turning a slow LibreOffice process into an opaque 500.
                raise DocumentExtractTimeout("旧版 Word 转换超时，请使用较小或较新的文档") from exc
            converted = path / "source.docx"
            if proc.returncode or not converted.exists():
                raise ValueError("旧版 Word 转换失败，请核对文件是否损坏")
            converted_data = converted.read_bytes()
            _validate_docx_archive(converted_data)
            text = _docx(converted_data)
        warnings.append("旧版 Word 经 LibreOffice 转换，请核对表格与版式内容")
    elif suffix == ".docx":
        try:
            _validate_docx_archive(content)
            text = _docx(content)
        except Exception as exc:
            raise ValueError("Word 文档损坏或无法读取") from exc
    else:
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("加密 PDF 暂不支持")
            if len(reader.pages) > 100:
                raise ValueError("PDF 超过 100 页限制")
            renderer = None
            pages = []
            recognized = False
            try:
                for index, page in enumerate(reader.pages):
                    page_text = (page.extract_text() or "").strip()
                    if len(page_text) < 20:
                        renderer = renderer or pdfium.PdfDocument(content)
                        image = renderer[index].render(scale=2).to_pil()
                        buffer = io.BytesIO()
                        image.save(buffer, format="PNG")
                        ocr_text, warning = _ocr(buffer.getvalue())
                        if warning:
                            warnings.append(f"第 {index + 1} 页：{warning}")
                        if ocr_text:
                            page_text = ocr_text
                            warnings.append(f"第 {index + 1} 页使用 OCR，请核对识别内容")
                    recognized = recognized or bool(page_text)
                    pages.append(f"[第 {index + 1} 页]\n{page_text}")
                text = "\n\n".join(pages) if recognized else ""
            finally:
                # Pdfium owns native handles outside Python's GC. Close the
                # renderer after OCR pages so repeated uploads cannot exhaust
                # the API process over time.
                close = getattr(renderer, "close", None)
                if callable(close):
                    close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 损坏或无法读取") from exc
    if MAX_EXTRACTED_TEXT_CHARS and len(text) > MAX_EXTRACTED_TEXT_CHARS:
        text = text[:MAX_EXTRACTED_TEXT_CHARS]
        warnings.append(
            f"可分析文本超过 {MAX_EXTRACTED_TEXT_CHARS:,} 个字符，已按块截取前段内容；原始文件仍可下载"
        )
    if not text.strip():
        warnings.append("未提取到可分析文字，不能启动 AI 分析")
    return text, warnings
