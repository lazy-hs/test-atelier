import io
import json
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["AI_STUDIO_SESSION_SECRET"] = "test-only-session-secret-32-characters-long"
os.environ["AI_STUDIO_ADMIN_PASSWORD"] = "test-admin-password-123"

from api.db import Base, engine  # noqa: E402
from api.importers import parse  # noqa: E402
from api.main import app  # noqa: E402


@pytest.fixture
def client():
    Base.metadata.drop_all(engine)
    with TestClient(app) as client:
        yield client


def auth(client):
    token = client.get("/api/v1/auth/csrf").json()["token"]
    result = client.post("/api/v1/auth/login", json={"username": "admin", "password": "test-admin-password-123"}, headers={"x-csrf-token": token})
    assert result.status_code == 200, result.text
    return {"x-csrf-token": result.json()["csrf"]}


def workbook():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["测试用例ID", "用例名称", "用例等级", "是否冒烟用例", "用例步骤"])
    sheet.append(["001", "正常用例", "1", None, "=HYPERLINK(\"x\")"])
    sheet.append(["002", "原始否", "3", "否", "普通步骤"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_import_review_export_and_csrf(client):
    headers = auth(client)
    project = client.post("/api/v1/projects", json={"name": "监控中心"}, headers=headers).json()
    prefix = f"/api/v1/projects/{project['id']}"
    denied = client.post(f"{prefix}/imports/preview", files={"file": ("cases.xlsx", workbook())})
    assert denied.status_code == 403
    preview = client.post(f"{prefix}/imports/preview", files={"file": ("cases.xlsx", workbook())}, headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()["count"] == 2
    batch_id = preview.json()["id"]
    assert client.post(f"{prefix}/imports/{batch_id}/confirm", headers=headers).json()["count"] == 2
    assert client.post(f"{prefix}/imports/{batch_id}/confirm", headers=headers).json()["count"] == 2
    cases = client.get(f"{prefix}/cases").json()
    assert cases["total"] == 2
    assert cases["items"][0]["original_id"] == "001"
    assert cases["items"][0]["smoke"] == ""
    case_id = cases["items"][0]["id"]
    payload = {"version": 1, "final_level": "1", "final_smoke": "", "reason": "人工核对"}
    reviewed = client.put(f"{prefix}/cases/{case_id}/review", json=payload, headers=headers)
    assert reviewed.status_code == 200, reviewed.text
    overview = client.get(f"{prefix}/overview").json()
    assert overview["case_count"] == 2
    assert overview["pending_count"] == 1
    assert overview["import_count"] == 1
    assert overview["recent_activity"][0]["kind"] == "import"
    assert client.put(f"{prefix}/cases/{case_id}/review", json=payload, headers=headers).status_code == 409
    export = client.get(f"{prefix}/export")
    assert export.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
    assert sheet.max_row == 3
    assert sheet.cell(2, 1).value == "001"
    assert sheet.cell(2, 8).value == "已审核"
    assert sheet.cell(3, 8).value == "待审核"
    assert sheet.cell(1, 11).value == "负责人"


def test_duplicate_import_is_atomic(client):
    headers = auth(client)
    project_id = client.post("/api/v1/projects", json={"name": "P"}, headers=headers).json()["id"]
    prefix = f"/api/v1/projects/{project_id}"
    for expected in (200, 409):
        batch = client.post(f"{prefix}/imports/preview", files={"file": ("cases.xlsx", workbook())}, headers=headers).json()
        response = client.post(f"{prefix}/imports/{batch['id']}/confirm", headers=headers)
        assert response.status_code == expected
    assert client.get(f"{prefix}/cases").json()["total"] == 2


def test_real_xls_sample_if_present():
    sample = Path("/Users/lazy/Downloads/UGOS_PRO_测试用例_20261015141029.xls")
    if not sample.exists():
        pytest.skip("验收样本不在本机")
    headers, rows, warnings = parse(sample.name, sample.read_bytes())
    assert len(headers) == 13
    assert len(rows) == 100
    assert len({row[0] for row in rows}) == 100
    assert rows[0][0] == "001252040"
    assert sum(row[11] is None for row in rows) == 78
    assert sum(row[11] == "否" for row in rows) == 22
    assert warnings


@pytest.mark.parametrize("format_name", ["content.json", "content.xml"])
def test_xmind_leaf_paths(format_name):
    content = (
        json.dumps([{"rootTopic": {"title": "监控中心", "children": {"attached": [
            {"title": "播放", "children": {"attached": [{"title": "正常播放"}]}}
        ]}}}], ensure_ascii=False).encode()
        if format_name.endswith("json") else
        b'<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0"><sheet><topic><title>Root</title><children><topics type="attached"><topic><title>Leaf</title></topic></topics></children></topic></sheet></xmap-content>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(format_name, content)
    headers, rows, warnings = parse("cases.xmind", buffer.getvalue())
    assert headers == ["测试用例ID", "用例名称", "用例目录"]
    assert len(rows) == 1
    assert rows[0][1] == ("正常播放" if format_name.endswith("json") else "Leaf")
    assert rows[0][2] == ("监控中心 / 播放" if format_name.endswith("json") else "Root")
    assert warnings


def test_reader_cannot_import_or_review(client):
    headers = auth(client)
    project_id = client.post("/api/v1/projects", json={"name": "P"}, headers=headers).json()["id"]
    user = client.post("/api/v1/users", json={"username": "reader", "password": "reader-password-123", "role": "reader"}, headers=headers).json()
    client.post(f"/api/v1/projects/{project_id}/members", json={"user_id": user["id"]}, headers=headers)
    client.post("/api/v1/auth/logout", headers=headers)
    token = client.get("/api/v1/auth/csrf").json()["token"]
    reader = client.post("/api/v1/auth/login", json={"username": "reader", "password": "reader-password-123"}, headers={"x-csrf-token": token}).json()
    response = client.post(f"/api/v1/projects/{project_id}/imports/preview", files={"file": ("cases.xlsx", workbook())}, headers={"x-csrf-token": reader["csrf"]})
    assert response.status_code == 403
    assert client.get(f"/api/v1/projects/{project_id}/cases").status_code == 200


def test_formatted_numeric_id_keeps_display_zeros():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["测试用例ID", "用例名称"])
    sheet.append([42, "零前缀"])
    sheet["A2"].number_format = "000000"
    buffer = io.BytesIO()
    book.save(buffer)
    _, rows, _ = parse("numbers.xlsx", buffer.getvalue())
    assert rows[0][0] == "000042"
