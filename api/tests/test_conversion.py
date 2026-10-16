import io
import os
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["AI_STUDIO_SESSION_SECRET"] = "test-only-session-secret-32-characters-long"
os.environ["AI_STUDIO_ADMIN_PASSWORD"] = "test-admin-password-123"

from api.db import Base, engine  # noqa: E402
from api.main import app  # noqa: E402
from api import conversion_routes  # noqa: E402
from api.conversion.service import preview as convert_preview  # noqa: E402


@pytest.fixture
def context(tmp_path, monkeypatch):
    monkeypatch.setattr(conversion_routes, "DATA_DIR", tmp_path / "conversions")
    Base.metadata.drop_all(engine)
    with TestClient(app) as client:
        token = client.get("/api/v1/auth/csrf").json()["token"]
        logged = client.post("/api/v1/auth/login", json={"username": "admin", "password": "test-admin-password-123"}, headers={"x-csrf-token": token})
        headers = {"x-csrf-token": logged.json()["csrf"]}
        project = client.post("/api/v1/projects", json={"name": "P"}, headers=headers).json()
        yield client, headers, f"/api/v1/projects/{project['id']}"


def source_file():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["用例目录", "用例名称", "用例等级", "前置条件", "用例步骤", "预期结果", "是否是冒烟用例"])
    sheet.append(["登录", "账号登录", "2", "账号存在", "输入账号", "进入首页", None])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def sample_xmind():
    xml = b'<xmap-content xmlns="urn:xmind:xmap:xmlns:content:2.0"><sheet><topic><title>Root</title><children><topics type="attached"><topic><title>Login</title><children><topics type="attached"><topic><title>Works</title><marker-refs><marker-ref marker-id="priority-2"/></marker-refs><children><topics type="attached"><topic><title>Step 1</title><children><topics type="attached"><topic><title>Expected</title></topic></topics></children></topic></topics></children></topic></topics></children></topic></topics></children></topic></sheet></xmap-content>'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("content.xml", xml)
    return buffer.getvalue()


def form(mode, output="after"):
    return {"conversion_mode": mode, "output_mode": output, "prefix_path": "迭代", "case_type": "业务功能用例", "case_status": "正常"}


@pytest.mark.parametrize("mode,output,filename,content,target_kind,columns", [
    ("excel_to_excel", "after", "cases.xlsx", source_file(), "table", 15),
    ("excel_to_excel", "after", "cases.csv", "用例名称,前置条件,用例步骤,预期结果\nA,B,C,D\n".encode(), "table", 15),
    ("xmind_to_excel", "before", "cases.xmind", sample_xmind(), "table", 7),
    ("xmind_to_final_excel", "after", "cases.xmind", sample_xmind(), "table", 15),
    ("excel_to_xmind", "after", "cases.xlsx", source_file(), "tree", 0),
])
def test_conversion_modes_and_history(context, mode, output, filename, content, target_kind, columns):
    client, headers, root = context
    url = root + "/conversion"
    args = {"data": form(mode, output), "files": {"file": (filename, content)}, "headers": headers}
    assert client.post(url + "/preview", **{**args, "headers": {}}).status_code == 403
    result = client.post(url + "/preview", **args)
    assert result.status_code == 200, result.text
    assert result.json()["target_kind"] == target_kind
    assert len(result.json()["target_headers"]) == columns
    response = client.post(url + "/export", **args)
    assert response.status_code == 200, response.text
    if target_kind == "tree":
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert "content.xml" in archive.namelist()
    else:
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert book.active.max_column == columns
    history = client.get(url + "/history").json()
    assert len(history["records"]) == 1
    overview = client.get(root + "/overview").json()
    assert overview["conversion_count"] == 1
    assert any(item["kind"] == "conversion" for item in overview["recent_activity"])
    record_id = history["records"][0]["id"]
    assert client.get(url + f"/history/{record_id}").json()["preview"]["row_count"] == 1
    assert client.get(url + f"/history/{record_id}/download").content == response.content
    assert client.delete(url + f"/history/{record_id}", headers=headers).status_code == 200
    assert client.get(url + "/history").json()["records"] == []


def test_project_config_and_isolated_history(context):
    client, headers, root = context
    url = root + "/conversion"
    result = client.patch(url + "/config", json={"default_prefix_path": "另一个目录", "history_retention_days": 30}, headers=headers)
    assert result.status_code == 200
    assert result.json()["default_prefix_path"] == "另一个目录"
    assert client.get(url + "/config").json()["history_retention_days"] == 30
    assert client.patch(url + "/config", json={"history_retention_days": 4}, headers=headers).status_code == 422
    other = client.post("/api/v1/projects", json={"name": "Other"}, headers=headers).json()["id"]
    assert client.get(f"/api/v1/projects/{other}/conversion/config").json()["default_prefix_path"] == ""
    response = client.post(url + "/export", files={"file": ("cases.xlsx", source_file())}, data=form("excel_to_excel"), headers=headers)
    assert response.status_code == 200
    assert client.get(f"/api/v1/projects/{other}/conversion/history").json()["records"] == []
    assert client.delete(url + "/history", headers=headers).json()["deleted_count"] == 1


def test_first_reads_do_not_write_and_reader_permissions(context):
    client, headers, root = context
    url = root + "/conversion"
    responses = [client.get(url + "/config"), client.get(url + "/history")]
    assert [response.status_code for response in responses] == [200, 200]
    from api.db import ConversionSettings, SessionLocal
    with SessionLocal() as db:
        assert db.get(ConversionSettings, int(root.split("/")[-1])) is None
    assert client.delete(url + "/history").status_code == 403
    user = client.post("/api/v1/users", json={"username": "reader", "password": "reader-password-123", "role": "reader"}, headers=headers).json()
    project_id = int(root.split("/")[-1])
    client.post(f"{root}/members", json={"user_id": user["id"]}, headers=headers)
    other = client.post("/api/v1/projects", json={"name": "Private"}, headers=headers).json()["id"]
    client.post("/api/v1/auth/logout", headers=headers)
    token = client.get("/api/v1/auth/csrf").json()["token"]
    logged = client.post("/api/v1/auth/login", json={"username": "reader", "password": "reader-password-123"}, headers={"x-csrf-token": token}).json()
    reader_headers = {"x-csrf-token": logged["csrf"]}
    assert project_id != other
    assert client.get(url + "/history").status_code == 200
    assert client.get(f"/api/v1/projects/{other}/conversion/history").status_code == 403
    assert client.patch(url + "/config", json={"default_prefix_path": "other"}, headers=reader_headers).status_code == 403
    assert client.post(url + "/preview", data=form("excel_to_excel"), files={"file": ("cases.xlsx", source_file())}, headers=reader_headers).status_code == 403


def test_xmind_bridge_does_not_populate_internal_smoke():
    result = convert_preview(sample_xmind(), "a.xmind", "xmind_to_excel", "before", "", "", "")
    assert result["row_count"] == 1
    assert result["target_rows"][0]["用例等级"] == "2"


def test_excel_formula_text_is_escaped_and_xmind_expansion_limited():
    from api.conversion.converter import export_source_workbook
    from api.conversion.xmind_converter import parse_xmind
    book = openpyxl.load_workbook(io.BytesIO(export_source_workbook([{"用例名称": "=HYPERLINK(\"bad\")"}])))
    assert book.active.cell(2, 2).data_type == "s"
    assert book.active.cell(2, 2).value.startswith("'")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("content.xml", b"<x/>")
        archive.writestr("large.bin", b"a" * (81 * 1024 * 1024))
    with pytest.raises(ValueError, match="安全限制"):
        parse_xmind(buffer.getvalue(), "large.xmind", " - ")
