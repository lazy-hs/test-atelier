import io
import os
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["AI_STUDIO_SESSION_SECRET"] = "test-only-session-secret-32-characters-long"
os.environ["AI_STUDIO_ADMIN_PASSWORD"] = "test-admin-password-123"

from api import ai_engine, classification, suggestion_jobs  # noqa: E402
from api.db import Base, engine  # noqa: E402
from api.main import app  # noqa: E402


@pytest.fixture
def context(monkeypatch):
    monkeypatch.setattr(suggestion_jobs, "enqueue", suggestion_jobs.process_job)
    Base.metadata.drop_all(engine)
    with TestClient(app) as client:
        token = client.get("/api/v1/auth/csrf").json()["token"]
        response = client.post("/api/v1/auth/login", json={"username": "admin", "password": "test-admin-password-123"},
                               headers={"x-csrf-token": token})
        headers = {"x-csrf-token": response.json()["csrf"]}
        project = client.post("/api/v1/projects", json={"name": "分类项目"}, headers=headers).json()["id"]
        book = openpyxl.Workbook(); sheet = book.active
        sheet.append(["测试用例ID", "用例名称", "用例等级", "是否冒烟用例", "用例目录", "用例类型", "端类型"])
        sheet.append(["001", "账号登录", "1", None, "账号 - 登录", "业务功能用例", "Web"])
        sheet.append(["002", "异常登录", "3", "否", "其他", "异常场景用例", "Web"])
        stream = io.BytesIO(); book.save(stream)
        root = f"/api/v1/projects/{project}"
        preview = client.post(f"{root}/imports/preview", files={"file": ("cases.xlsx", stream.getvalue())}, headers=headers).json()
        assert client.post(f"{root}/imports/{preview['id']}/confirm", headers=headers).status_code == 200
        ids = [row["id"] for row in client.get(f"{root}/cases").json()["items"]]
        yield client, headers, root, ids


def publish_taxonomy(client, headers, root):
    draft = client.post(f"{root}/taxonomy/versions", headers=headers).json()
    definitions = {
        "business_module": [{"name": "账号", "description": "账号能力"}],
        "business_scenario": [{"name": "登录", "description": "登录流程"}],
        "test_dimension": [{"name": "业务功能用例", "description": "功能检查"}],
        "endpoint_type": [{"name": "Web", "description": "浏览器端"}],
    }
    saved = client.put(f"{root}/taxonomy/versions/{draft['id']}", json={"revision": draft["revision"], "definitions": definitions}, headers=headers)
    assert saved.status_code == 200, saved.text
    version = saved.json()
    assert client.post(f"{root}/taxonomy/versions/{draft['id']}/publish", json={"revision": draft["revision"]}, headers=headers).status_code == 409
    published = client.post(f"{root}/taxonomy/versions/{draft['id']}/publish", json={"revision": version["revision"]}, headers=headers)
    assert published.status_code == 200, published.text
    return published.json()


def wait_job(client, root, job_id):
    for _ in range(100):
        job = client.get(f"{root}/suggestion-jobs/{job_id}").json()
        if job["status"] in ("done", "failed", "partial"):
            return job
        time.sleep(.025)
    pytest.fail("建议任务未完成")


def test_taxonomy_versions_and_rule_candidates_preserve_unknowns(context):
    client, headers, root, ids = context
    initial = client.post(f"{root}/suggestions/rules", json={"case_ids": [ids[0]]}, headers=headers).json()["suggestions"][0]
    assert initial["level"] == "1" and initial["smoke"] is None
    assert set(initial["categories"].values()) == {None}
    assert client.get(f"{root}/cases/{ids[0]}").json()["final_smoke"] is None
    draft = client.post(f"{root}/taxonomy/versions", headers=headers).json()
    assert draft["status"] == "draft"
    assert client.post(f"{root}/taxonomy/versions", headers=headers).status_code == 409
    assert client.post(f"{root}/taxonomy/versions/{draft['id']}/publish", json={"revision": 1}, headers=headers).status_code == 422
    duplicate = {key: [{"name": "相同"}, {"name": "相同"}] for key in classification.DIMENSIONS}
    assert client.put(f"{root}/taxonomy/versions/{draft['id']}", json={"revision": 1, "definitions": duplicate}, headers=headers).status_code == 422
    definitions = {key: [{"name": name, "description": ""}] for key, name in zip(classification.DIMENSIONS,
                   ("账号", "登录", "业务功能用例", "Web"))}
    saved = client.put(f"{root}/taxonomy/versions/{draft['id']}", json={"revision": 1, "definitions": definitions}, headers=headers).json()
    assert client.put(f"{root}/taxonomy/versions/{draft['id']}", json={"revision": 1, "definitions": definitions}, headers=headers).status_code == 409
    published = client.post(f"{root}/taxonomy/versions/{draft['id']}/publish", json={"revision": saved["revision"]}, headers=headers).json()
    assert client.put(f"{root}/taxonomy/versions/{draft['id']}", json={"revision": published["revision"], "definitions": definitions}, headers=headers).status_code == 409
    proposal = client.post(f"{root}/suggestions/rules", json={"case_ids": [ids[0]]}, headers=headers).json()["suggestions"][0]
    assert proposal["taxonomy_version_id"] == published["id"]
    assert set(proposal["category_labels"].values()) == {"账号", "登录", "业务功能用例", "Web"}
    assert proposal["smoke"] is None
    assert client.post(f"{root}/suggestions/rules", json={"case_ids": [ids[0]]}, headers=headers).json()["suggestions"][0]["id"] == proposal["id"]
    another = client.post(f"{root}/taxonomy/versions", headers=headers).json()
    assert another["number"] == 2 and another["definitions"] == published["definitions"]


def test_model_suggestions_decisions_audit_and_export(context, monkeypatch):
    client, headers, root, ids = context
    version = publish_taxonomy(client, headers, root)
    rule = client.post(f"{root}/suggestions/rules", json={"case_ids": [ids[0], ids[1]]}, headers=headers).json()["suggestions"]
    identifiers = {key: values[0]["id"] for key, values in version["definitions"].items()}

    def model(_prompt, _schema, _provider):
        return {"categories": identifiers, "level": "2", "smoke": "是", "rationale": "需要人工确认核心链路",
                "questions": ["请确认冒烟时间"]}

    monkeypatch.setattr(ai_engine, "request_structured", model)
    created = client.post(f"{root}/suggestions/model", json={"case_ids": [ids[0]], "provider": "local"}, headers=headers)
    assert created.status_code == 200, created.text
    job = wait_job(client, root, created.json()["id"])
    assert job["status"] == "done" and job["completed"] == 1
    suggestions = client.get(f"{root}/cases/{ids[0]}/suggestions").json()
    model_suggestion = next(item for item in suggestions if item["source"] == "model")
    assert model_suggestion["categories"] == identifiers
    assert client.get(f"{root}/cases?status=pending&suggestion=pending").json()["total"] == 2
    assert client.post(f"{root}/suggestions/{model_suggestion['id']}/decision", json={"action": "accept", "case_version": 1, "reason": " "}, headers=headers).status_code == 422
    response = client.post(f"{root}/suggestions/{model_suggestion['id']}/decision", json={
        "action": "accept", "case_version": 1, "reason": "人工修改冒烟判断",
        "choice": {"categories": identifiers, "level": "2", "smoke": "否"}}, headers=headers)
    assert response.status_code == 200, response.text
    case = response.json()["case"]
    assert case["original_id"] == "001" and case["smoke"] == "" and case["final_smoke"] == "否"
    assert case["final_level"] == "2" and case["final_categories"] == identifiers
    assert case["final_taxonomy_version_id"] == version["id"] and case["version"] == 2
    accepted = next(item for item in client.get(f"{root}/cases/{ids[0]}/suggestions").json() if item["id"] == model_suggestion["id"])
    assert accepted["status"] == "accepted" and accepted["stale"] is False
    assert client.post(f"{root}/suggestions/{model_suggestion['id']}/decision", json={"action": "reject", "case_version": 2, "reason": "重复"}, headers=headers).status_code == 409
    assert client.get(f"{root}/cases?status=pending&suggestion=pending").json()["total"] == 1
    assert client.get(f"{root}/cases/{ids[0]}/activity").json()["events"][0]["action"] == "suggestion"
    rejected = client.post(f"{root}/suggestions/{rule[1]['id']}/decision", json={"action": "reject", "case_version": 1, "reason": "原始等级需要重看"}, headers=headers)
    assert rejected.status_code == 200
    assert client.get(f"{root}/cases/{ids[1]}").json()["final_level"] is None
    book = openpyxl.load_workbook(io.BytesIO(client.get(f"{root}/export").content))
    headers_row = [cell.value for cell in book.active[1]]
    assert "最终业务模块" in headers_row and "建议依据" in headers_row
    assert book.active.cell(2, headers_row.index("最终业务模块") + 1).value == "账号"
    assert client.post(f"{root}/suggestions/{rule[0]['id']}/decision", json={"action": "accept", "case_version": 1, "reason": "过期"}, headers=headers).status_code == 409
    other = client.post("/api/v1/projects", json={"name": "另一个项目"}, headers=headers).json()["id"]
    assert client.get(f"/api/v1/projects/{other}/cases/{ids[0]}/suggestions").status_code == 404
    assert client.post(f"/api/v1/projects/{other}/suggestions/{model_suggestion['id']}/decision", json={"action": "reject", "case_version": 2, "reason": "越权"}, headers=headers).status_code == 404


def test_invalid_model_ids_fail_without_writing_case(context, monkeypatch):
    client, headers, root, ids = context
    publish_taxonomy(client, headers, root)
    monkeypatch.setattr(ai_engine, "request_structured", lambda *_args: {
        "categories": {key: "invented" for key in classification.DIMENSIONS},
        "level": "1", "smoke": "否", "rationale": "未经验证", "questions": []})
    job_id = client.post(f"{root}/suggestions/model", json={"case_ids": [ids[0]]}, headers=headers).json()["id"]
    assert wait_job(client, root, job_id)["status"] == "failed"
    assert client.get(f"{root}/cases/{ids[0]}/suggestions").json() == []
    assert client.get(f"{root}/cases/{ids[0]}").json()["final_smoke"] is None


def test_partial_model_retry_is_idempotent(context, monkeypatch):
    client, headers, root, ids = context
    version = publish_taxonomy(client, headers, root)
    identifiers = {key: values[0]["id"] for key, values in version["definitions"].items()}
    def sometimes_invalid(prompt, _schema, _provider):
        categories = {key: "invented" for key in identifiers} if "异常登录" in prompt else identifiers
        return {"categories": categories, "level": "2", "smoke": None, "rationale": "仅供人工审核", "questions": []}
    monkeypatch.setattr(ai_engine, "request_structured", sometimes_invalid)
    job_id = client.post(f"{root}/suggestions/model", json={"case_ids": ids}, headers=headers).json()["id"]
    partial = wait_job(client, root, job_id)
    assert partial["status"] == "partial" and partial["completed"] == 2 and len(partial["errors"]) == 1
    assert len(client.get(f"{root}/cases/{ids[0]}/suggestions").json()) == 1
    assert client.get(f"{root}/cases/{ids[1]}/suggestions").json() == []
    monkeypatch.setattr(ai_engine, "request_structured", lambda *_args: {
        "categories": identifiers, "level": "3", "smoke": None, "rationale": "复核后重试", "questions": []})
    assert client.post(f"{root}/suggestion-jobs/{job_id}/retry", headers=headers).status_code == 200
    retried = wait_job(client, root, job_id)
    assert retried["status"] == "done" and retried["completed"] == 2 and retried["errors"] == []
    assert len(client.get(f"{root}/cases/{ids[0]}/suggestions").json()) == 1
    assert len(client.get(f"{root}/cases/{ids[1]}/suggestions").json()) == 1


def test_assigned_reviewer_cannot_decide_other_members_suggestion(context):
    client, headers, root, ids = context
    first = client.post("/api/v1/users", json={"username": "alice", "password": "alice-password-123", "role": "reviewer"}, headers=headers).json()
    second = client.post("/api/v1/users", json={"username": "bob", "password": "bob-password-123", "role": "reviewer"}, headers=headers).json()
    for member in (first, second):
        client.post(f"{root}/members", json={"user_id": member["id"]}, headers=headers)
    assigned = client.put(f"{root}/cases/{ids[0]}/assignment", json={"version": 1, "assignee_id": first["id"]}, headers=headers)
    assert assigned.status_code == 200
    proposal = client.post(f"{root}/suggestions/rules", json={"case_ids": [ids[0]]}, headers=headers).json()["suggestions"][0]
    client.post("/api/v1/auth/logout", headers=headers)
    csrf = client.get("/api/v1/auth/csrf").json()["token"]
    logged = client.post("/api/v1/auth/login", json={"username": "bob", "password": "bob-password-123"},
                         headers={"x-csrf-token": csrf}).json()
    response = client.post(f"{root}/suggestions/{proposal['id']}/decision", json={"action": "accept", "case_version": 2, "reason": "越权"},
                           headers={"x-csrf-token": logged["csrf"]})
    assert response.status_code == 403
    assert client.get(f"{root}/cases/{ids[0]}").json()["final_level"] is None


def test_reader_can_only_read_published_taxonomy_and_suggestions(context):
    client, headers, root, ids = context
    draft = client.post(f"{root}/taxonomy/versions", headers=headers).json()
    reader = client.post("/api/v1/users", json={"username": "reader", "password": "reader-password-123", "role": "reader"}, headers=headers).json()
    client.post(f"{root}/members", json={"user_id": reader["id"]}, headers=headers)
    client.post("/api/v1/auth/logout", headers=headers)
    csrf = client.get("/api/v1/auth/csrf").json()["token"]
    logged = client.post("/api/v1/auth/login", json={"username": "reader", "password": "reader-password-123"}, headers={"x-csrf-token": csrf}).json()
    auth = {"x-csrf-token": logged["csrf"]}
    assert client.get(f"{root}/taxonomy").json()["draft"] is None
    assert client.get(f"{root}/taxonomy/versions/{draft['id']}").status_code == 404
    assert client.post(f"{root}/taxonomy/versions", headers=auth).status_code == 403
    assert client.post(f"{root}/suggestions/rules", json={"case_ids": ids}, headers=auth).status_code == 403
    assert client.post(f"{root}/suggestions/model", json={"case_ids": ids}, headers=auth).status_code == 403
