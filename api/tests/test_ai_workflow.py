import io
import os
import re
import time
import zipfile

import openpyxl
import pytest
from docx import Document
from PIL import Image
from pypdf import PdfWriter
from fastapi.testclient import TestClient

os.environ["DATABASE_URL"] = "sqlite://"
os.environ["AI_STUDIO_SESSION_SECRET"] = "test-only-session-secret-32-characters-long"
os.environ["AI_STUDIO_ADMIN_PASSWORD"] = "test-admin-password-123"

from api import ai_engine, ai_jobs, ai_routes, skill_runner  # noqa: E402
from api import main as api_main  # noqa: E402
from api.db import AiJob, AiQueueState, Base, SessionLocal, engine  # noqa: E402
from api.main import app  # noqa: E402


@pytest.fixture
def context(tmp_path, monkeypatch):
    monkeypatch.setattr(ai_routes, "DOCUMENT_DIR", tmp_path / "documents")
    Base.metadata.drop_all(engine)
    with TestClient(app) as client:
        csrf = client.get("/api/v1/auth/csrf").json()["token"]
        login = client.post("/api/v1/auth/login", json={"username": "admin", "password": "test-admin-password-123"}, headers={"x-csrf-token": csrf}).json()
        headers = {"x-csrf-token": login["csrf"]}
        project = client.post("/api/v1/projects", json={"name": "P"}, headers=headers).json()
        yield client, headers, f"/api/v1/projects/{project['id']}"


def wait_job(client, root, job_id):
    for _ in range(80):
        result = client.get(f"{root}/ai/jobs/{job_id}").json()
        if result["status"] in {"done", "failed"}:
            return result
        time.sleep(.025)
    pytest.fail(f"AI job did not finish: {result}")


def fake_model(prompt, schema, _provider):
    if "requirements" in schema["properties"]:
        source = re.search(r"文档 ID: ([a-f0-9]+)", prompt)
        return {"requirements": [{"id": "REQ-1", "title": "登录", "summary": "账号登录", "source_ids": [source.group(1)],
                                   "acceptance": ["正确账号进入首页"], "methods": ["EP"]}],
                "ambiguities": [{"id": "AQ-1", "description": "失败次数限制", "critical": True, "decision": ""}], "risks": ["连续失败"]}
    return {"cases": [{"module": ["登录"], "title": "有效账号登录", "priority": "P1", "precondition": "有账号",
                       "steps": [{"action": "输入正确账号", "expected": "进入首页"}], "requirement_ids": ["REQ-1"],
                       "method": "EP", "label": "Web", "remark": "", "smoke": ""}]}


def test_document_analysis_authoring_and_exports(context, monkeypatch, tmp_path):
    client, headers, root = context
    live = client.get("/api/v1/health/live", headers={"x-request-id": "workflow-trace"})
    assert live.status_code == 200
    assert live.headers["x-request-id"] == "workflow-trace"
    assert live.headers["x-frame-options"] == "DENY"
    metrics = client.get("/api/v1/metrics")
    assert metrics.status_code == 200
    assert "test_atelier_requests_total" in metrics.text
    monkeypatch.setenv("METRICS_TOKEN", "metrics-test-token")
    metric_events = []
    with monkeypatch.context() as local:
        local.setattr(api_main.logger, "warning", lambda message, **kwargs: metric_events.append((message, kwargs)))
        assert client.get("/api/v1/metrics").status_code == 401
    assert metric_events and metric_events[0][0] == "metrics.auth.failed"
    assert client.get("/api/v1/metrics", headers={"X-Metrics-Token": "metrics-test-token"}).status_code == 200
    monkeypatch.delenv("METRICS_TOKEN")
    assert client.get("/api/v1/health/ready").json()["status"] == "ready"
    assert client.get("/api/v1/health/worker").json()["status"] == "disabled"
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", True)
    with SessionLocal() as db:
        queue_state = db.get(AiQueueState, 1)
        queue_state.active_jobs = 7
        db.commit()
    heartbeat = tmp_path / "worker-heartbeat"
    heartbeat.write_text('{"timestamp": 1, "release": "worker-release-test"}')
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT", str(heartbeat))
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT_SECONDS", "20")
    worker_health = client.get("/api/v1/health/worker").json()
    assert worker_health["status"] == "ready"
    assert worker_health["worker_release"] == "worker-release-test"
    worker_metrics = client.get("/api/v1/metrics").text
    assert "test_atelier_ai_jobs_active 7" in worker_metrics
    assert "test_atelier_worker_up 1" in worker_metrics
    assert "test_atelier_worker_heartbeat_threshold_seconds 60.000" in worker_metrics
    old = time.time() - 120
    os.utime(heartbeat, (old, old))
    stale = client.get("/api/v1/health/worker")
    assert stale.status_code == 503 and "过期" in stale.json()["detail"]
    stale_metrics = client.get("/api/v1/metrics").text
    assert "test_atelier_worker_up 0" in stale_metrics
    heartbeat.unlink()
    assert client.get("/api/v1/health/worker").status_code == 503
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", False)
    monkeypatch.setenv("TEST_ATELIER_SKILL_ROOT", "/path/that/does/not/exist")
    assert client.get("/api/v1/health/ready").status_code == 503
    monkeypatch.delenv("TEST_ATELIER_SKILL_ROOT")
    monkeypatch.setattr(ai_engine, "request_structured", fake_model)
    assert client.post(f"{root}/documents", files={"file": ("r.md", b"# Login\nCorrect account opens home")}).status_code == 403
    upload = client.post(f"{root}/documents", files={"file": ("r.md", b"# Login\nCorrect account opens home")}, headers=headers)
    assert upload.status_code == 200, upload.text
    doc_id = upload.json()["id"]
    assert client.get(f"{root}/documents/{doc_id}").json()["text"].startswith("# Login")
    assert client.get(f"{root}/documents/{doc_id}/download").content.startswith(b"# Login")
    response = client.post(f"{root}/analyses", json={"source_ids": [doc_id]},
                           headers={**headers, "x-request-id": "analysis-enqueue"})
    assert response.status_code == 200, response.text
    analysis_id = response.json()["analysis_id"]
    analysis_job_id = response.json()["job_id"]
    assert client.get(f"{root}/ai/jobs/{analysis_job_id}").json()["request_id"] == "analysis-enqueue"
    assert wait_job(client, root, analysis_job_id)["status"] == "done"
    analysis_prompt = client.get(f"{root}/ai/jobs/{analysis_job_id}/prompt")
    assert analysis_prompt.status_code == 200
    assert analysis_prompt.json()["prompt_calls"] == 1
    assert "按需求分析 Skill" in analysis_prompt.json()["prompts"][0]
    assert analysis_prompt.json()["prompt_length"] == len(analysis_prompt.json()["prompts"][0])
    structured_prompt = client.get(f"{root}/ai/jobs/{analysis_job_id}/prompts").json()
    assert structured_prompt["prompts"][0]["index"] == 1
    assert structured_prompt["prompts"][0]["provider"] == "local"
    assert structured_prompt["prompts"][0]["text"] == analysis_prompt.json()["prompts"][0]
    report = client.get(f"{root}/analyses/{analysis_id}").json()
    assert report["result"]["requirements"][0]["source_ids"] == [doc_id]
    assert client.post(f"{root}/analyses/{analysis_id}/generate", headers=headers).status_code == 409
    edited = client.patch(f"{root}/analyses/{analysis_id}/report", json={"version": report["version"], "report_markdown": "# 人工修订\n登录规则"}, headers=headers)
    assert edited.status_code == 200
    report = edited.json()
    decision = client.patch(f"{root}/analyses/{analysis_id}/ambiguities/AQ-1", json={"version": report["version"], "decision": "5 次锁定"}, headers=headers)
    assert decision.status_code == 200, decision.text
    assert decision.json()["result"]["report_markdown"] == "# 人工修订\n登录规则"
    assert client.patch(f"{root}/analyses/{analysis_id}/ambiguities/AQ-1", json={"version": report["version"], "decision": "8 次"}, headers=headers).status_code == 409
    generated = client.post(f"{root}/analyses/{analysis_id}/generate", headers=headers)
    assert generated.status_code == 200, generated.text
    cases_job_id = generated.json()["job_id"]
    assert wait_job(client, root, cases_job_id)["status"] == "done"
    cases_prompt = client.get(f"{root}/ai/jobs/{cases_job_id}/prompt")
    assert cases_prompt.status_code == 200
    assert cases_prompt.json()["prompt_calls"] == 1
    assert "按 test-cases Skill" in cases_prompt.json()["prompts"][0]
    drafts = client.get(f"{root}/drafts?analysis_id={analysis_id}").json()
    assert len(drafts) == 1
    draft = drafts[0]
    assert draft["payload"]["smoke"] == ""
    update = client.patch(f"{root}/drafts/{draft['id']}", json={"version": draft["version"], "payload": {**draft["payload"], "title": "人工确认登录"}}, headers=headers)
    assert update.status_code == 200, update.text
    xmind = client.get(f"{root}/drafts/export?analysis_id={analysis_id}&format=xmind")
    assert xmind.status_code == 200, xmind.text
    with zipfile.ZipFile(io.BytesIO(xmind.content)) as archive:
        assert "content.xml" in archive.namelist()
    excel = client.get(f"{root}/drafts/export?analysis_id={analysis_id}&format=xlsx")
    assert excel.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    assert sheet.cell(2, 3).value == "人工确认登录"
    published = client.post(f"{root}/drafts/{draft['id']}/publish", headers=headers)
    assert published.status_code == 200, published.text
    cases = client.get(f"{root}/cases").json()["items"]
    assert cases[0]["smoke"] == ""
    assert cases[0]["level"] == "2"
    assert client.post(f"{root}/drafts/{draft['id']}/publish", headers=headers).json()["case_id"] == published.json()["case_id"]
    overview = client.get(f"{root}/overview").json()
    assert overview["case_count"] == 1
    assert overview["analysis_count"] == 1
    assert overview["document_count"] == 1
    assert overview["draft_count"] == 0
    assert {item["kind"] for item in overview["recent_activity"]} >= {"analysis", "draft", "document"}


def test_readiness_rejects_traffic_during_shutdown(context, monkeypatch):
    client, _headers, _root = context
    monkeypatch.setattr(api_main, "_ready_for_traffic", False)
    response = client.get("/api/v1/health/ready")
    assert response.status_code == 503
    assert "启动或停止" in response.json()["detail"]


def test_readiness_failure_logging_is_rate_limited(monkeypatch):
    records = []
    monkeypatch.setattr(api_main, "_readiness_log_last", {})
    monkeypatch.setattr(api_main.logger, "log", lambda level, message, **kwargs: records.append((level, message, kwargs)))
    api_main._log_readiness_failure("worker.readiness.stale", component="external-worker")
    api_main._log_readiness_failure("worker.readiness.stale", component="external-worker")
    assert len(records) == 1
    assert records[0][1] == "worker.readiness.stale"


def test_worker_heartbeat_defaults_under_data_directory(monkeypatch, tmp_path):
    monkeypatch.delenv("AI_STUDIO_WORKER_HEARTBEAT", raising=False)
    monkeypatch.setenv("AI_STUDIO_DATA_DIR", str(tmp_path / "runtime-data"))
    assert api_main._worker_heartbeat_path() == str(tmp_path / "runtime-data" / ".worker-heartbeat")


@pytest.mark.skipif(os.name == "nt", reason="symlink behavior is not portable on Windows")
def test_worker_health_rejects_symlinked_heartbeat(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", True)
    target = tmp_path / "outside-heartbeat"
    target.write_text("now", encoding="ascii")
    linked = tmp_path / "worker-heartbeat"
    linked.symlink_to(target)
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT", str(linked))
    response = client.get("/api/v1/health/worker")
    assert response.status_code == 503


def test_worker_health_treats_metadata_race_as_unavailable(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", True)
    heartbeat = tmp_path / "worker-heartbeat"
    heartbeat.write_text('{"timestamp": 1, "release": "worker-build"}', encoding="utf-8")
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT", str(heartbeat))
    monkeypatch.setattr(api_main, "_heartbeat_release", lambda _path: (_ for _ in ()).throw(OSError("replaced")))

    response = client.get("/api/v1/health/worker")

    assert response.status_code == 503
    assert "元数据不可用" in response.json()["detail"]


def test_worker_health_logs_release_mismatch(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", True)
    monkeypatch.setenv("AI_STUDIO_RELEASE", "api-build-42")
    monkeypatch.setattr(api_main, "_readiness_log_last", {})
    records = []
    monkeypatch.setattr(api_main.logger, "log", lambda level, message, **kwargs: records.append((level, message, kwargs)))
    heartbeat = tmp_path / "worker-heartbeat"
    heartbeat.write_text('{"timestamp": 1, "release": "worker-build-41"}', encoding="utf-8")
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT", str(heartbeat))

    response = client.get("/api/v1/health/worker")

    assert response.status_code == 200
    assert response.json()["worker_release"] == "worker-build-41"
    assert any(message == "worker.readiness.version_mismatch"
               and fields.get("extra", {}).get("error_type") == "ReleaseMismatch"
               for _level, message, fields in records)


def test_worker_health_logs_unknown_release_for_legacy_worker(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    monkeypatch.setattr(ai_jobs, "EXTERNAL_WORKER", True)
    monkeypatch.setenv("AI_STUDIO_RELEASE", "api-build-42")
    monkeypatch.setattr(api_main, "_readiness_log_last", {})
    records = []
    monkeypatch.setattr(api_main.logger, "log", lambda level, message, **kwargs: records.append((level, message, kwargs)))
    heartbeat = tmp_path / "worker-heartbeat"
    heartbeat.write_text("legacy-heartbeat", encoding="utf-8")
    monkeypatch.setenv("AI_STUDIO_WORKER_HEARTBEAT", str(heartbeat))

    response = client.get("/api/v1/health/worker")

    assert response.status_code == 200
    assert "worker_release" not in response.json()
    assert any(message == "worker.readiness.version_unknown"
               and fields.get("extra", {}).get("error_type") == "MissingWorkerRelease"
               for _level, message, fields in records)


def test_readiness_rejects_unwritable_data_directory(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    monkeypatch.setenv("AI_STUDIO_DATA_DIR", str(tmp_path / "missing-data"))
    response = client.get("/api/v1/health/ready")
    assert response.status_code == 503
    assert "数据目录不可写" in response.json()["detail"]


@pytest.mark.skipif(os.name == "nt", reason="symlink behavior is not portable on Windows")
def test_readiness_rejects_symlinked_data_directory(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    real_data = tmp_path / "real-data"
    real_data.mkdir()
    linked_data = tmp_path / "data-link"
    linked_data.symlink_to(real_data, target_is_directory=True)
    monkeypatch.setenv("AI_STUDIO_DATA_DIR", str(linked_data))
    response = client.get("/api/v1/health/ready")
    assert response.status_code == 503
    assert "数据目录不可写" in response.json()["detail"]


@pytest.mark.skipif(os.name == "nt", reason="symlink behavior is not portable on Windows")
def test_document_download_rejects_replaced_symlink(context, tmp_path):
    client, headers, root = context
    uploaded = client.post(f"{root}/documents", files={"file": ("secret.txt", b"private")}, headers=headers)
    assert uploaded.status_code == 200
    document_id = uploaded.json()["id"]
    artifact = ai_routes.DOCUMENT_DIR / root.rsplit("/", 1)[-1] / f"{document_id}.bin"
    outside = tmp_path / "outside.txt"
    outside.write_text("outside", encoding="utf-8")
    artifact.unlink()
    artifact.symlink_to(outside)
    response = client.get(f"{root}/documents/{document_id}/download")
    assert response.status_code == 404


def test_readiness_rejects_incomplete_configured_skill_root(context, monkeypatch, tmp_path):
    client, _headers, _root = context
    skill_root = tmp_path / "qa-doc-skills"
    (skill_root / "skills" / "hub").mkdir(parents=True)
    (skill_root / "AGENTS.md").write_text("# rules", encoding="utf-8")
    (skill_root / "skills" / "hub" / "SKILL.md").write_text("# hub", encoding="utf-8")
    monkeypatch.setenv("TEST_ATELIER_SKILL_ROOT", str(skill_root))
    response = client.get("/api/v1/health/ready")
    assert response.status_code == 503
    assert "工作流文件" in response.json()["detail"]


def test_ai_queue_full_returns_retry_after(context, monkeypatch):
    client, headers, root = context
    upload = client.post(f"{root}/documents", files={"file": ("r.md", b"# Login\nCorrect account opens home")}, headers=headers)
    assert upload.status_code == 200
    monkeypatch.setattr(ai_jobs, "enqueue", lambda *_args, **_kwargs: False)
    response = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers)
    assert response.status_code == 503
    assert response.headers["retry-after"] == "5"


def test_word_extraction_and_offline_failure(context, monkeypatch):
    client, headers, root = context
    word = Document(); word.add_paragraph("需求：登录后展示首页")
    buffer = io.BytesIO(); word.save(buffer)
    upload = client.post(f"{root}/documents", files={"file": ("req.docx", buffer.getvalue())}, headers=headers)
    assert "登录后展示首页" in upload.json()["text"]
    def unavailable(*_args):
        raise RuntimeError("模型离线")
    monkeypatch.setattr(ai_engine, "request_structured", unavailable)
    job = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers).json()
    assert "模型离线" in wait_job(client, root, job["job_id"])["error"]
    assert client.get(f"{root}/cases").status_code == 200
    assert client.post(f"{root}/ai/jobs/{job['job_id']}/retry", headers=headers).status_code == 200
    assert wait_job(client, root, job["job_id"])["status"] == "failed"
    assert client.get(f"{root}/ai/jobs/{job['job_id']}/prompt").json()["prompt_calls"] == 2


def test_failed_ai_job_keeps_prompt_for_diagnostics(context, monkeypatch):
    client, headers, root = context

    def unavailable(*_args):
        raise RuntimeError("模型离线")

    monkeypatch.setattr(ai_engine, "request_structured", unavailable)
    upload = client.post(f"{root}/documents", files={"file": ("req.md", "登录需求".encode())}, headers=headers)
    job = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers).json()
    assert wait_job(client, root, job["job_id"])["status"] == "failed"
    prompt = client.get(f"{root}/ai/jobs/{job['job_id']}/prompt").json()
    assert prompt["status"] == "failed"
    assert prompt["prompt_calls"] == 1
    assert "文档 ID:" in prompt["prompts"][0]


def test_prompt_viewer_normalizes_legacy_structured_entries(context):
    client, headers, root = context
    upload = client.post(f"{root}/documents", files={"file": ("req.md", b"login")}, headers=headers)
    job = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers).json()
    wait_job(client, root, job["job_id"])
    with SessionLocal() as db:
        row = db.get(AiJob, job["job_id"])
        row.prompts = [{"text": "legacy prompt"}, {"unexpected": "value"}, "current prompt"]
        row.prompt_calls = 3
        row.prompt_length = 0
        db.commit()
    legacy = client.get(f"{root}/ai/jobs/{job['job_id']}/prompt")
    assert legacy.status_code == 200
    assert legacy.json()["prompts"] == ["legacy prompt", "current prompt"]
    structured = client.get(f"{root}/ai/jobs/{job['job_id']}/prompts")
    assert structured.status_code == 200
    assert [item["text"] for item in structured.json()["prompts"]] == ["legacy prompt", "current prompt"]


def test_ai_job_error_redacts_provider_credentials(context, monkeypatch):
    client, headers, root = context

    def unavailable(*_args):
        raise RuntimeError("provider failed postgresql://app:db-secret@db/app password=hunter2")

    monkeypatch.setattr(ai_engine, "request_structured", unavailable)
    upload = client.post(f"{root}/documents", files={"file": ("req.md", "登录需求".encode())}, headers=headers)
    job = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers).json()
    result = wait_job(client, root, job["job_id"])
    assert result["status"] == "failed"
    assert "db-secret" not in result["error"]
    assert "hunter2" not in result["error"]
    assert "[REDACTED]" in result["error"]


def test_skill_failure_does_not_expose_host_path_in_job_error(context, monkeypatch):
    client, headers, root = context

    def unavailable(*_args):
        raise skill_runner.SkillRunnerError("技能文件不存在：/Users/private/qa-doc-skills/AGENTS.md")

    monkeypatch.setattr(skill_runner, "analyze", unavailable)
    upload = client.post(f"{root}/documents", files={"file": ("req.md", "登录需求".encode())}, headers=headers)
    job = client.post(f"{root}/analyses", json={"source_ids": [upload.json()["id"]]}, headers=headers).json()
    result = wait_job(client, root, job["job_id"])
    assert result["status"] == "failed"
    assert result["error"] == "技能工作流不可用，请检查技能仓配置"
    assert "/Users/private" not in result["error"]


def test_image_and_scanned_pdf_require_real_text(context, monkeypatch):
    from api import document_extract
    client, headers, root = context
    image = io.BytesIO(); Image.new("RGB", (200, 80), "white").save(image, format="PNG")
    monkeypatch.setattr(document_extract, "_ocr", lambda _data: ("", "OCR 不可用"))
    picture = client.post(f"{root}/documents", files={"file": ("screen.png", image.getvalue())}, headers=headers).json()
    assert picture["text"] == ""
    assert client.post(f"{root}/analyses", json={"source_ids": [picture["id"]]}, headers=headers).status_code == 422
    writer = PdfWriter(); writer.add_blank_page(width=200, height=100)
    pdf = io.BytesIO(); writer.write(pdf)
    blank = client.post(f"{root}/documents", files={"file": ("scan.pdf", pdf.getvalue())}, headers=headers).json()
    assert blank["text"] == ""
    monkeypatch.setattr(document_extract, "_ocr", lambda _data: ("登录按钮", None))
    scanned = client.post(f"{root}/documents", files={"file": ("scan.pdf", pdf.getvalue())}, headers=headers).json()
    assert "登录按钮" in scanned["text"]
    assert any("使用 OCR" in warning for warning in scanned["warnings"])


def test_document_project_isolation_and_reader_access(context):
    client, headers, root = context
    uploaded = client.post(f"{root}/documents", files={"file": ("r.md", b"# Private")}, headers=headers).json()
    other = client.post("/api/v1/projects", json={"name": "Other"}, headers=headers).json()["id"]
    assert client.get(f"/api/v1/projects/{other}/documents/{uploaded['id']}").status_code == 404
    assert client.post(f"{root}/analyses", json={"source_ids": [uploaded["id"]]}, headers={}).status_code == 403
    user = client.post("/api/v1/users", json={"username": "reader", "password": "reader-password-123", "role": "reader"}, headers=headers).json()
    client.post(f"{root}/members", json={"user_id": user["id"]}, headers=headers)
    client.post("/api/v1/auth/logout", headers=headers)
    token = client.get("/api/v1/auth/csrf").json()["token"]
    login = client.post("/api/v1/auth/login", json={"username": "reader", "password": "reader-password-123"}, headers={"x-csrf-token": token}).json()
    reader_headers = {"x-csrf-token": login["csrf"]}
    assert client.get(f"{root}/documents/{uploaded['id']}").status_code == 200
    assert client.get(f"/api/v1/projects/{other}/documents").status_code == 403
    assert client.post(f"{root}/documents", files={"file": ("r.md", b"# No")}, headers=reader_headers).status_code == 403
    assert client.post(f"{root}/analyses", json={"source_ids": [uploaded["id"]]}, headers=reader_headers).status_code == 403


def test_multidocument_dedup_and_quality_gate(monkeypatch):
    def model(prompt, schema, _provider):
        ids = re.findall(r'"source_ids": \["([a-f0-9]+)"\]', prompt)
        if "文档 ID:" in prompt:
            ids = re.findall(r"文档 ID: ([a-f0-9]+)", prompt)
        return {"requirements": [{"id": "REQ-1", "title": "登录", "summary": "登录成功", "source_ids": ids,
                                   "acceptance": ["展示首页"], "methods": ["EP"]}], "ambiguities": [], "risks": []}
    monkeypatch.setattr(ai_engine, "request_structured", model)
    documents = [{"id": "a" * 32, "filename": "a.md", "text": "登录成功展示首页"},
                 {"id": "b" * 32, "filename": "b.md", "text": "有效账号可进入首页"}]
    result = ai_engine.analyze(documents, "local")
    assert len(result["requirements"]) == 1
    assert set(result["requirements"][0]["source_ids"]) == {"a" * 32, "b" * 32}
    duplicate = [{"module": ["登录"], "title": "同名", "requirement_ids": ["REQ-1"], "method": "EP",
                  "steps": [{"action": "操作", "expected": "结果"}]}] * 2
    assert "存在重复的模块/用例标题" in ai_engine.assess_cases(duplicate, result)["errors"]
