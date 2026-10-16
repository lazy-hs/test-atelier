import csv
import io
import re
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE_HEADERS = [
    "用例目录",
    "用例名称",
    "用例等级",
    "前置条件",
    "用例步骤",
    "预期结果",
    "是否是冒烟用例",
]

TARGET_HEADERS = [
    "用例目录",
    "用例名称",
    "需求ID",
    "前置条件",
    "用例步骤",
    "预期结果",
    "用例类型",
    "用例状态",
    "用例等级",
    "创建人",
    "备注",
    "是否冒烟用例",
    "端类型",
    "是否已经接口自动化",
    "是否已经UI自动化",
]

HEADER_ALIASES = {
    "是否是冒烟用例": ("是否是冒烟用例", "是否冒烟用例"),
}

REQUIRED_HEADERS = ("用例名称", "前置条件", "用例步骤", "预期结果")


class ConversionError(ValueError):
    pass


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _safe_excel_text(value: Any) -> Any:
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", _clean_text(value))


def _resolve_header(headers: list[str], canonical: str) -> int | None:
    aliases = HEADER_ALIASES.get(canonical, (canonical,))
    normalized = {_normalize_header(header): index for index, header in enumerate(headers)}
    for alias in aliases:
        if _normalize_header(alias) in normalized:
            return normalized[_normalize_header(alias)]
    return None


def _read_xlsx(content: bytes) -> tuple[str, list[str], list[list[Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ConversionError("无法读取该 Excel 文件，请确认文件未损坏且格式为 .xlsx。") from exc

    best_match: tuple[int, str, list[str], list[list[Any]]] | None = None
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            continue
        headers = [_clean_text(value) for value in values[0]]
        score = sum(_resolve_header(headers, header) is not None for header in SOURCE_HEADERS)
        rows = [list(row) for row in values[1:]]
        if best_match is None or score > best_match[0]:
            best_match = (score, sheet.title, headers, rows)

    if best_match is None:
        raise ConversionError("Excel 中没有可读取的工作表。")
    return best_match[1], best_match[2], best_match[3]


def _read_csv(content: bytes) -> tuple[str, list[str], list[list[Any]]]:
    decoded = None
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ConversionError("无法识别 CSV 文件编码，请使用 UTF-8 或 GB18030。")
    reader = list(csv.reader(io.StringIO(decoded)))
    if not reader:
        raise ConversionError("CSV 文件为空。")
    return "CSV", reader[0], reader[1:]


def parse_source(content: bytes, filename: str) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        sheet_name, headers, rows = _read_xlsx(content)
    elif suffix == ".csv":
        sheet_name, headers, rows = _read_csv(content)
    else:
        raise ConversionError("当前支持 .xlsx 和 .csv 文件。")

    missing = [header for header in REQUIRED_HEADERS if _resolve_header(headers, header) is None]
    if missing:
        raise ConversionError(f"缺少必填列：{'、'.join(missing)}。")

    indexes = {header: _resolve_header(headers, header) for header in SOURCE_HEADERS}
    cases = []
    for source_row_number, row in enumerate(rows, start=2):
        values = {}
        for header, index in indexes.items():
            values[header] = _clean_text(row[index]) if index is not None and index < len(row) else ""
        if not any(values.values()):
            continue
        if not values["用例名称"]:
            raise ConversionError(f"第 {source_row_number} 行缺少用例名称。")
        values["source_row"] = source_row_number
        cases.append(values)

    if not cases:
        raise ConversionError("文件中没有可转换的用例数据。")

    return {"sheet_name": sheet_name, "headers": headers, "cases": cases}


def _join_path(prefix_path: str, module_path: str, separator: str) -> str:
    parts = [part.strip() for part in (prefix_path, module_path) if part and part.strip()]
    return separator.join(parts)


def transform_cases(source_cases: list[dict], prefix_path: str, case_type: str, case_status: str, separator: str) -> list[dict]:
    prefix_path = prefix_path.strip()
    if not prefix_path:
        raise ConversionError("请填写前置路径。")
    if not case_type:
        raise ConversionError("请选择用例类型。")

    converted = []
    for case in source_cases:
        converted.append(
            {
                "用例目录": _join_path(prefix_path, case.get("用例目录", ""), separator),
                "用例名称": case.get("用例名称", ""),
                "需求ID": "",
                "前置条件": case.get("前置条件", ""),
                "用例步骤": case.get("用例步骤", ""),
                "预期结果": case.get("预期结果", ""),
                "用例类型": case_type,
                "用例状态": case_status,
                "用例等级": case.get("用例等级", ""),
                "创建人": "",
                "备注": "",
                "是否冒烟用例": case.get("是否是冒烟用例", ""),
                "端类型": "",
                "是否已经接口自动化": "",
                "是否已经UI自动化": "",
            }
        )
    return converted


def build_preview(content: bytes, filename: str, prefix_path: str, case_type: str, case_status: str, separator: str) -> dict:
    parsed = parse_source(content, filename)
    converted = transform_cases(parsed["cases"], prefix_path, case_type, case_status, separator)
    original = [{header: case.get(header, "") for header in SOURCE_HEADERS} for case in parsed["cases"]]
    return {
        "filename": filename,
        "sheet_name": parsed["sheet_name"],
        "row_count": len(converted),
        "original_headers": SOURCE_HEADERS,
        "target_headers": TARGET_HEADERS,
        "original": original,
        "converted": converted,
        "preview_type": "excel_to_excel",
        "source_kind": "table",
        "source_headers": SOURCE_HEADERS,
        "source_rows": original,
        "target_kind": "table",
        "target_headers": TARGET_HEADERS,
        "target_rows": converted,
        "export_label": "导出 Excel",
    }


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def export_workbook(converted: list[dict], template_path: Path) -> bytes:
    if not template_path.exists():
        raise ConversionError(f"找不到导入模板：{template_path.name}")

    workbook = load_workbook(template_path)
    sheet = workbook.active
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    header_style_sources = [sheet.cell(1, column) for column in range(1, len(TARGET_HEADERS) + 1)]
    for column, header in enumerate(TARGET_HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        _copy_cell_style(header_style_sources[column - 1], cell)

    body_border = Border(bottom=Side(style="hair", color="D9DEE8"))
    for row_index, case in enumerate(converted, start=2):
        for column_index, header in enumerate(TARGET_HEADERS, start=1):
            cell = sheet.cell(row_index, column_index, _safe_excel_text(case.get(header, "")))
            cell.font = Font(name="微软雅黑", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = body_border
        if row_index % 2 == 0:
            for column_index in range(1, len(TARGET_HEADERS) + 1):
                sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="F8FAFC")
        sheet.row_dimensions[row_index].height = 52

    widths = [48, 34, 16, 44, 44, 48, 18, 14, 12, 16, 18, 18, 18, 24, 22]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:O{max(2, len(converted) + 1)}"
    sheet.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_source_workbook(cases: list[dict], sheet_name: str = "测试用例") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31] or "测试用例"

    header_fill = PatternFill("solid", fgColor="EDEDFC")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="2F3650")
    body_font = Font(name="微软雅黑", size=10, color="1F2937")
    border = Border(bottom=Side(style="hair", color="D9DEE8"))

    for column_index, header in enumerate(SOURCE_HEADERS, start=1):
        cell = sheet.cell(1, column_index, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, case in enumerate(cases, start=2):
        for column_index, header in enumerate(SOURCE_HEADERS, start=1):
            cell = sheet.cell(row_index, column_index, _safe_excel_text(case.get(header, "")))
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        if row_index % 2 == 0:
            for column_index in range(1, len(SOURCE_HEADERS) + 1):
                sheet.cell(row_index, column_index).fill = PatternFill("solid", fgColor="F8FAFC")
        sheet.row_dimensions[row_index].height = 52

    widths = [42, 36, 12, 44, 44, 48, 18]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{max(2, len(cases) + 1)}"
    sheet.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_output_name(source_filename: str) -> str:
    stem = Path(source_filename).stem
    stem = re.sub(r"[\\/:*?\"<>|]", "_", stem).strip(" .") or "测试用例"
    return f"{stem}_格式转换结果.xlsx"
